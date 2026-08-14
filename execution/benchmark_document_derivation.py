"""Benchmark reproducible de derivación y contexto recuperado de documentos.

Este gate usa un corpus sintético local y deliberadamente pequeño. Mide lo que
realmente se enviaría tras seleccionar ``top_k`` (no compara solo el Markdown
derivado con el extractor anterior), además de preservación y recall de evidencia.

La selección léxica es un *proxy offline* determinista del RRF de producción: no
requiere Postgres, pgvector ni una API de embeddings. Por eso los resultados solo
describen estos fixtures y NO demuestran por sí solos una reducción de 25 % en
expedientes reales.
"""
from __future__ import annotations

import io
import importlib.util
import json
import re
import statistics
import sys
import time
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from mia.ingest.document_derivation import derive_office_document  # noqa: E402
from mia.memory.tokens import estimate_tokens  # noqa: E402

TOP_K = 3
N_SECTIONS = 24
RUNS = 5


def fixture() -> tuple[bytes, list[dict[str, str]]]:
    """Construye un DOCX con evidencia trazable y bastante ruido estructural."""
    from docx import Document

    document = Document()
    cases: list[dict[str, str]] = []
    for index in range(N_SECTIONS):
        marker = f"EVIDENCIA-{index:02d}"
        keyword = f"terminounico{index:02d}"
        document.add_heading(f"Sección {index:02d}: {keyword}", 1)
        document.add_paragraph(
            f"Antecedentes de la sección {index:02d}. "
            + "Este párrafo de contexto no sustituye la fuente original. " * 3
        )
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Fuente", "Hallazgo"
        table.cell(1, 0).text = f"Anexo {index:02d}"
        table.cell(1, 1).text = f"{marker} {keyword}"
        cases.append({"query": f"¿Qué dice {keyword}?", "evidence": marker})
    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), cases


def _terms(text: str) -> set[str]:
    normalized = unicodedata.normalize("NFKD", text.casefold())
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return {term for term in re.findall(r"[a-z0-9]{3,}", normalized)
            if term not in {"que", "dice"}}


def retrieve_top_k(blocks: tuple[dict, ...], query: str, *, top_k: int = TOP_K) -> list[dict]:
    """Proxy léxico offline: devuelve exactamente el contexto seleccionado top-k."""
    query_terms = _terms(query)
    ranked = sorted(
        blocks,
        key=lambda block: (
            len(query_terms & _terms(str(block.get("text", "")))),
            -int(block.get("position", 0)),
        ),
        reverse=True,
    )
    return ranked[:top_k]


def run_benchmark() -> dict[str, int | float | str]:
    blob, cases = fixture()
    conversion_ms: list[float] = []
    derived = None
    for _ in range(RUNS):
        started = time.perf_counter()
        derived = derive_office_document("fixture.docx", blob)
        conversion_ms.append((time.perf_counter() - started) * 1000)
    if derived is None:
        raise RuntimeError("AnyDoc no está disponible; instala el entorno fijado del backend")

    preserved = sum(case["evidence"] in derived.markdown for case in cases)
    full_tokens = estimate_tokens(derived.markdown)
    retrieval_ms: list[float] = []
    retrieved_tokens: list[int] = []
    hits = 0
    for case in cases:
        started = time.perf_counter()
        selected = retrieve_top_k(derived.blocks, case["query"])
        retrieval_ms.append((time.perf_counter() - started) * 1000)
        context = "\n\n".join(str(block.get("text", "")) for block in selected)
        retrieved_tokens.append(estimate_tokens(context))
        hits += int(case["evidence"] in context)

    spec = importlib.util.find_spec("anydoc")
    package_dir = Path(spec.origin).parent if spec and spec.origin else None
    package_size = (sum(path.stat().st_size for path in package_dir.rglob("*") if path.is_file())
                    if package_dir and package_dir.exists() else 0)
    median_retrieved = int(statistics.median(retrieved_tokens))
    reduction = 1.0 - (median_retrieved / max(1, full_tokens))
    return {
        "corpus": "synthetic_docx_fixture",
        "queries": len(cases),
        "top_k": TOP_K,
        "chunks": len(derived.blocks),
        "full_context_tokens": full_tokens,
        "retrieved_context_tokens_median": median_retrieved,
        "context_reduction_fixture_pct": round(reduction * 100, 2),
        "structural_evidence_recall_pct": round(100 * preserved / len(cases), 2),
        "retrieval_evidence_recall_at_k_pct": round(100 * hits / len(cases), 2),
        "conversion_median_ms": round(statistics.median(conversion_ms), 2),
        "retrieval_median_ms": round(statistics.median(retrieval_ms), 4),
        "anydoc_payload_bytes": package_size,
        "retrieval_mode": "offline_lexical_proxy_not_production_rrf",
        "claim_scope": "fixture_only_no_25_percent_claim",
    }


def main() -> None:
    result = run_benchmark()
    assert result["structural_evidence_recall_pct"] == 100.0, result
    assert result["retrieval_evidence_recall_at_k_pct"] == 100.0, result
    assert result["retrieved_context_tokens_median"] < result["full_context_tokens"], result
    print(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
