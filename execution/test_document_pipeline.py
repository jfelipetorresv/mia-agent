"""
Mia · test_document_pipeline.py — gate de CP9 (equipo de especialistas de documentos).

Verifica OFFLINE (sin DB, sin LLM, sin red — todo con dobles):
  A. Instrucciones L8 de los especialistas nuevos (facts/research) + cruce (analysis).
  B. Cableado del grafo: intake → facts → research → analysis → draft → verification
     → hitl_checkpoint → finalize.
  C. Nodos facts/research/analysis: prompts correctos con _llm doblado.
  D. Especialista de verificación: escaneo/anotación determinista de citas.
  E. Emisión Word: draft_to_docx produce un .docx válido con el formato esperado.
  F. Superficie API: endpoint .docx registrado; verification expuesto (SSE + GET).

Exit 0 = PASS · 1 = FAIL.       .venv\\Scripts\\python.exe execution\\test_document_pipeline.py
"""
import asyncio
import inspect
import io
import sys
from pathlib import Path

from dotenv import load_dotenv

if sys.platform == "win32":
    asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())
    # La consola cp1252 no imprime '→' — el gate no debe caerse por un print.
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")
sys.path.insert(0, str(ROOT / "backend"))

_results: list = []


def check(name: str, ok) -> None:
    _results.append((name, bool(ok)))
    print(("  [OK]   " if ok else "  [FAIL] ") + name)


async def run_gate() -> None:
    from mia.agent import prompt_builder as pb
    from mia.agents import graph as graph_mod
    from mia.agents import research, verification
    from mia.agents.graph import MatterGraphBuilder
    from mia.output.docx_export import FOOTER_NOTE, draft_to_docx

    # ── A · instrucciones L8 de los especialistas ────────────────────────────
    check("cp9-01 · GRAPH_NODE_INSTRUCTIONS tiene los especialistas facts y research",
          "facts" in pb.GRAPH_NODE_INSTRUCTIONS and "research" in pb.GRAPH_NODE_INSTRUCTIONS)
    check("cp9-02 · el especialista de hechos NO analiza derecho (lo dice su instrucción)",
          "NO analices el derecho" in pb.GRAPH_NODE_INSTRUCTIONS["facts"])
    check("cp9-03 · el investigador exige [VERIFICAR] para lo que no venga del corpus",
          "[VERIFICAR]" in pb.GRAPH_NODE_INSTRUCTIONS["research"])
    check("cp9-04 · el cruce (analysis) conserva el bloque de cierre estructurado (CP5/CP6)",
          pb.DIAGNOSIS_CLOSING_HEADER in pb.GRAPH_NODE_INSTRUCTIONS["analysis"]
          and pb.DIAGNOSIS_CLOSING_FOOTER in pb.GRAPH_NODE_INSTRUCTIONS["analysis"]
          and "especialista de hechos" in pb.GRAPH_NODE_INSTRUCTIONS["analysis"])
    sys_facts = pb.build_graph_system({"soul_snapshot": {"content": "Voz sobria."}}, "facts")
    check("cp9-05 · build_graph_system('facts') compone las 10 capas (identidad+SOUL+L8)",
          pb.GRAPH_FALLBACK_IDENTITY in sys_facts and "Voz sobria." in sys_facts
          and pb.GRAPH_NODE_INSTRUCTIONS["facts"] in sys_facts)
    check("cp9-06 · build_graph_system('research') funciona y lleva la citación L3",
          pb.CITATION_POLICY in pb.build_graph_system({}, "research"))

    # ── B · cableado del grafo ───────────────────────────────────────────────
    builder = MatterGraphBuilder(trace_capture=_FakeTraceCapture())
    compiled = builder.build(checkpointer=None)
    drawable = compiled.get_graph()
    nodes = set(drawable.nodes)
    check("cp9-07 · el grafo tiene los 8 nodos del equipo",
          {"intake", "facts", "research", "analysis", "draft", "verification",
           "hitl_checkpoint", "finalize"} <= nodes)
    edges = {(e.source, e.target) for e in drawable.edges}
    expected = {("intake", "facts"), ("facts", "research"), ("research", "analysis"),
                ("analysis", "draft"), ("draft", "verification"),
                ("verification", "hitl_checkpoint"), ("hitl_checkpoint", "finalize")}
    check("cp9-08 · el orden del equipo es hechos→investigación→cruce→redacción→verificación",
          expected <= edges)

    # ── C · nodos con _llm doblado (sin LLM real) ────────────────────────────
    captured: dict = {}

    async def fake_llm(messages, *, task="main", state=None, md=None, shrink=None, node=""):
        captured[task + ":" + (md or {}).get("stage", "?")] = messages
        captured["last"] = messages
        return "SALIDA-DOBLADA", None

    builder._llm = fake_llm

    base_state = {
        "tenant_id": "00000000-0000-0000-0000-000000000001",
        "matter_id": "00000000-0000-0000-0000-000000000002",
        "messages": [{"role": "user", "content": "¿Operó la caducidad?"}],
        "documents": [{"id": "d1", "content": "El accidente ocurrió el 14 de marzo de 2024."}],
        "knowledge": [],
        "soul_snapshot": None,
        "metadata": {},
    }

    out = await builder.facts_node(dict(base_state))
    facts_md = out["metadata"]
    user_facts = captured["last"][1]["content"]
    check("cp9-09 · facts_node guarda los hechos en metadata['facts'] y stage='facts'",
          facts_md.get("facts") == "SALIDA-DOBLADA" and facts_md.get("stage") == "facts")
    check("cp9-10 · el prompt de hechos lleva la consulta y el expediente sellado (CP-S1)",
          "¿Operó la caducidad?" in user_facts and "<<<DOC 1>>>" in user_facts)
    check("cp9-11 · el system de hechos usa la instrucción del especialista",
          pb.GRAPH_NODE_INSTRUCTIONS["facts"] in captured["last"][0]["content"])

    # research_node con fuentes del corpus dobladas (sin DB)
    fake_sources = [{"tipo": "norma", "referencia": "Ley 640 de 2001",
                     "titulo": "Conciliación extrajudicial"}]

    async def fake_gather(tenant_id, query):
        captured["gather_query"] = query
        return ("<<<FUENTE 1 · Ley 640 de 2001>>>\nConciliación\n<<<FIN FUENTE 1>>>",
                fake_sources, ["co"])

    real_gather = research.gather_sources
    research.gather_sources = fake_gather
    try:
        st = dict(base_state); st["metadata"] = dict(facts_md)
        out = await builder.research_node(st)
    finally:
        research.gather_sources = real_gather
    res_md = out["metadata"]
    user_res = captured["last"][1]["content"]
    check("cp9-12 · research_node guarda memoria + fuentes + jurisdicciones en metadata",
          res_md.get("research") == "SALIDA-DOBLADA"
          and res_md.get("research_sources") == fake_sources
          and res_md.get("research_jurisdictions") == ["co"])
    check("cp9-13 · el prompt de investigación lleva hechos + fuentes fenceadas del corpus",
          "especialista de hechos" in user_res and "<<<FUENTE 1" in user_res)
    check("cp9-14 · la consulta FTS combina el mensaje del abogado con los hechos",
          "¿Operó la caducidad?" in captured["gather_query"])

    # research_node FAIL-SOFT: el corpus caído no tumba el turno
    async def broken_gather(tenant_id, query):
        raise RuntimeError("DB caída")
    research.gather_sources = graph_mod.research.gather_sources  # (mismo módulo)
    real_resolve = research.resolve_jurisdictions

    async def fake_resolve_fail(tenant_id):
        raise RuntimeError("DB caída")
    research.resolve_jurisdictions = fake_resolve_fail
    try:
        st = dict(base_state); st["metadata"] = dict(facts_md)
        out = await builder.research_node(st)
        check("cp9-15 · sin corpus (DB caída) la investigación sigue con [VERIFICAR] (fail-soft)",
              out["metadata"].get("research") == "SALIDA-DOBLADA"
              and out["metadata"].get("research_sources") == []
              and research.NO_SOURCES_NOTE in captured["last"][1]["content"])
    finally:
        research.resolve_jurisdictions = real_resolve

    # analysis_node (cruce): recibe el trabajo previo del equipo
    st = dict(base_state)
    st["metadata"] = {"facts": "1. Hecho probado.", "research": "Ley aplicable: X [VERIFICAR]."}
    out = await builder.analysis_node(st)
    user_an = captured["last"][1]["content"]
    check("cp9-16 · el cruce recibe hechos + memoria de investigación en su prompt",
          "1. Hecho probado." in user_an and "Ley aplicable: X" in user_an
          and "especialista de hechos" in user_an)
    st = dict(base_state); st["metadata"] = {}
    await builder.analysis_node(st)
    check("cp9-17 · sin facts/research (checkpoint viejo) el prompt del cruce queda como antes",
          "especialista de hechos" not in captured["last"][1]["content"])

    # ── D · especialista de verificación (determinista) ─────────────────────
    draft_txt = ("Se funda en el artículo 21 de la Ley 640 de 2001 y en la "
                 "Sentencia C-355 de 2006. También aplica la Ley 1437 de 2011 "
                 "[VERIFICAR texto vigente]. Ver Radicado 25000-23-41-000-2024.")
    annotated, report = verification.annotate_draft(
        draft_txt, sources=[{"tipo": "norma", "referencia": "Ley 640 de 2001"}])
    check("cp9-18 · detecta las citas del borrador (artículo+ley, sentencia, radicado)",
          report["citas"] >= 4)
    check("cp9-19 · la cita ya marcada [VERIFICAR] se respeta (no se duplica la marca)",
          report["marcadas"] >= 1 and "2011 [VERIFICAR] [VERIFICAR" not in annotated)
    check("cp9-20 · la cita respaldada en el corpus del turno NO se anota",
          report["respaldadas"] >= 1
          and "Ley 640 de 2001 [VERIFICAR]" not in annotated)
    check("cp9-21 · la sentencia SIN marca ni respaldo queda anotada [VERIFICAR]",
          report["anotadas"] >= 1 and "Sentencia C-355 de 2006 [VERIFICAR]" in annotated)
    check("cp9-22 · el informe trae el detalle por cita (estado de cada una)",
          isinstance(report["detalle"], list) and
          {d["estado"] for d in report["detalle"]} >= {"marcada", "respaldada", "anotada"})
    ann2, rep2 = verification.annotate_draft(annotated,
                                             sources=[{"tipo": "norma",
                                                       "referencia": "Ley 640 de 2001"}])
    check("cp9-23 · anotar dos veces es idempotente (no agrega marcas nuevas)",
          rep2["anotadas"] == 0 and ann2 == annotated)
    _, rep3 = verification.annotate_draft(
        "Sentencia T-406/92 sin marca.", extra_patterns=["\\bT-\\d{1,4}/\\d{2,4}"])
    check("cp9-24 · los patrones extra del pack de jurisdicción se aplican (horizontalidad)",
          rep3["citas"] >= 1 and rep3["anotadas"] >= 1)
    _, rep4 = verification.annotate_draft("texto", extra_patterns=["[regex-inválido"])
    check("cp9-25 · un patrón inválido del pack se ignora (fail-soft, no tumba el turno)",
          rep4["citas"] == 0)

    # verification_node dentro del grafo (patterns del pack doblados, sin DB)
    async def fake_patterns(tenant_id):
        return []
    real_patterns = research.citation_patterns_for
    research.citation_patterns_for = fake_patterns
    try:
        st = dict(base_state)
        st["draft"] = "Cita la Sentencia C-355 de 2006."
        st["metadata"] = {"research_sources": []}
        out = await builder.verification_node(st)
    finally:
        research.citation_patterns_for = real_patterns
    check("cp9-26 · verification_node anota el borrador y deja el informe en metadata",
          "[VERIFICAR]" in out["draft"] and out["metadata"]["verification"]["anotadas"] == 1)
    check("cp9-27 · el interrupt HITL expone el informe de verificación a la pantalla",
          '"verification"' in inspect.getsource(builder.hitl_checkpoint_node))

    # ── E · emisión Word ─────────────────────────────────────────────────────
    import docx as docx_lib
    md_draft = ("# CONTESTACIÓN DE LA DEMANDA\n\n---\n\n"
                "**SEÑOR JUEZ** del circuito.\n\n"
                "## I. HECHOS\n\n- Primero: ocurrió el accidente.\n"
                "1. Petición principal.\n\n"
                "Texto con **énfasis** y cita [VERIFICAR].")
    data = draft_to_docx(md_draft, title="Contestación", author="Lexia")
    d = docx_lib.Document(io.BytesIO(data))
    texts = [p.text for p in d.paragraphs]
    all_text = "\n".join(texts)
    check("cp9-28 · el .docx se genera y abre válido con el contenido del borrador",
          "SEÑOR JUEZ del circuito." in all_text and "Texto con énfasis y cita [VERIFICAR]."
          in all_text)
    check("cp9-29 · los adornos markdown NO llegan al Word (encabezados/negrilla/hr)",
          "**" not in all_text and "# " not in all_text and "---" not in all_text)
    heads = [p for p in d.paragraphs if p.style.name.startswith("Heading")]
    bold_runs = [r for p in d.paragraphs for r in p.runs if r.bold]
    check("cp9-30 · encabezados con estilo Heading y negrilla inline como runs bold",
          len(heads) >= 2 and any(r.text == "énfasis" for r in bold_runs))
    footer_text = d.sections[0].footer.paragraphs[0].text
    check("cp9-31 · el pie de página recuerda no radicar sin verificar citas",
          footer_text == FOOTER_NOTE and "[VERIFICAR]" in FOOTER_NOTE)
    check("cp9-32 · las propiedades del documento llevan título y autor",
          d.core_properties.title == "Contestación" and d.core_properties.author == "Lexia")

    # ── F · superficie API ───────────────────────────────────────────────────
    from mia.api.routes import stream as stream_mod
    from mia.api.routes import ux as ux_mod
    paths = {r.path for r in ux_mod.router.routes}
    check("cp9-33 · endpoint de descarga Word registrado (/api/matters/{id}/draft.docx)",
          "/api/matters/{matter_id}/draft.docx" in paths)
    src_ux = inspect.getsource(ux_mod.get_draft)
    check("cp9-34 · GET /draft expone el informe de verificación a la Pantalla 3",
          '"verification"' in src_ux)
    src_stream = inspect.getsource(stream_mod)
    check("cp9-35 · el SSE reenvía verification y anuncia los pasos del equipo (§G)",
          "verification=v.get(\"verification\")" in src_stream
          and "investigando normas y jurisprudencia" in src_stream)
    from mia.agents.context_recovery import NODE_BUDGET_FRACTION
    check("cp9-36 · los nodos nuevos tienen presupuesto de recorte (CP1 intacto)",
          "facts" in NODE_BUDGET_FRACTION and "research" in NODE_BUDGET_FRACTION)

    # ── G · hallazgos de la revisión capa 2 (M1 y M2) — no deben regresar ────
    # M1 · cupo de compresión POR NODO: que facts haya comprimido no le quita el
    # rescate a analysis (con cupo global, el turno moría en expedientes grandes).
    from mia.agent.error_classifier import LLMErrorKind
    from mia.agent.turn_llm_state import TurnLLMState
    turn = TurnLLMState()
    turn.mark_compressed("facts")
    check("cp9-37 · M1: el cupo de compresión es por nodo (facts no bloquea a analysis)",
          not turn.should_compress(LLMErrorKind.CONTEXT_TOO_LONG, "facts")
          and turn.should_compress(LLMErrorKind.CONTEXT_TOO_LONG, "analysis"))
    old = TurnLLMState.from_dict({"compression_attempted": True})
    check("cp9-38 · M1: metadata pre-CP9 (cupo global gastado) se respeta al deserializar",
          old.compression_attempted
          and not old.should_compress(LLMErrorKind.CONTEXT_TOO_LONG))
    rt = TurnLLMState.from_dict(turn.to_dict())
    check("cp9-39 · M1: el dict serializado conserva compression_attempted (gates CP1/H.5)",
          turn.to_dict()["compression_attempted"] is True
          and rt.compressed_stages == ["facts"])

    # M2 · ReDoS: el patrón de artículos no debe explotar con texto adversarial.
    import time as _time
    adversarial = "artículo 1" + "y1" * 40 + "!"
    t0 = _time.perf_counter()
    verification.annotate_draft(adversarial)
    elapsed = _time.perf_counter() - t0
    check(f"cp9-40 · M2: texto adversarial escanea en tiempo lineal ({elapsed*1000:.0f} ms < 1s)",
          elapsed < 1.0)
    _, rep_list = verification.annotate_draft(
        "Ver artículos 21, 22 y 23 de la Ley 640 de 2001 en la contestación.")
    check("cp9-41 · M2: la lista de artículos ('21, 22 y 23 de la Ley…') sigue detectándose",
          rep_list["citas"] >= 1 and "Ley 640 de 2001" in rep_list["detalle"][0]["cita"])


class _FakeTraceCapture:
    def capture(self, **kw):
        class _T:
            timestamp = "t0"
            def to_dict(self):
                return {}
        return _T()


if __name__ == "__main__":
    asyncio.run(run_gate())
    passed = sum(1 for _, ok in _results if ok)
    print(f"\n{passed}/{len(_results)} checks PASS")
    if all(ok for _, ok in _results):
        print("document pipeline OK — CP9 verificado (equipo de especialistas: hechos → "
              "investigación → cruce → redacción → verificación → Word).")
        sys.exit(0)
    sys.exit(1)
