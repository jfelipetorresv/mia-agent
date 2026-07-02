"""Mia · output.docx_export — borrador (markdown ligero) → documento Word (CP9).

El especialista de EMISIÓN convierte el borrador aprobado/pendiente en un .docx con
formato profesional de escrito jurídico. Determinista, sin LLM: el ESTILO DE VOZ del
usuario ya viene en el texto (lo puso el redactor vía SOUL/perfil/playbooks); aquí se
aplica el formato tipográfico del documento.

Soporta el markdown que emiten los nodos del grafo:
  #/##/### encabezados · **negrilla** inline · - viñetas · 1. listas numeradas ·
  --- (separador, se omite) · | tablas | (se aplanan a texto: v1) · ``` fences (las
  marcas ``` se omiten; su contenido se conserva literal)

El pie de página recuerda la regla de la casa: nada se radica sin que el abogado
verifique las citas marcadas [VERIFICAR] (§G: lenguaje del oficio, sin jerga).
"""
from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

FOOTER_NOTE = (
    "Borrador preparado por Mia para revisión del abogado responsable — "
    "no radicar sin verificar las citas marcadas [VERIFICAR]."
)

_BOLD_SPLIT = re.compile(r"\*\*(.+?)\*\*", re.DOTALL)
_BULLET = re.compile(r"^\s*[-*]\s+")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+")
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_HR = re.compile(r"^\s*(-{3,}|_{3,}|\*{3,})\s*$")
_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_TABLE_SEP = re.compile(r"^\s*\|[\s:|-]+\|\s*$")

BODY_FONT = "Arial"
BODY_SIZE_PT = 12


def _add_runs(paragraph, text: str) -> None:
    """Texto con **negrilla** inline → runs alternados (los ** no llegan al Word)."""
    parts = _BOLD_SPLIT.split(text)
    # split alterna: [normal, negrilla, normal, negrilla, ...]
    for i, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = i % 2 == 1


def _clean_inline(text: str) -> str:
    """Quita adornos que no deben llegar al documento (fences, backticks)."""
    return text.replace("`", "").strip()


def draft_to_docx(
    draft: str,
    *,
    title: str | None = None,
    author: str | None = None,
) -> bytes:
    """Convierte el borrador a un .docx (bytes) con formato de escrito jurídico.

    `title`/`author` van a las propiedades del documento (no al cuerpo: el cuerpo
    es el borrador tal cual lo redactó Mia y lo revisó el abogado).
    """
    doc = Document()

    # Tipografía base del oficio (cuerpo justificado, Arial 12).
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE_PT)

    props = doc.core_properties
    props.title = title or "Borrador"
    props.author = author or "Mia"

    in_fence = False
    for raw_line in (draft or "").splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            # el contenido del fence se conserva LITERAL (sin parsear markdown).
            doc.add_paragraph(line)
            continue
        if not stripped:
            continue
        if _HR.match(stripped):
            continue
        if _TABLE_SEP.match(stripped):
            continue
        if _TABLE_ROW.match(stripped):
            # v1: la fila de tabla se aplana a una línea de texto legible.
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            p = doc.add_paragraph()
            _add_runs(p, " · ".join(c for c in cells if c))
            continue

        m = _HEADING.match(stripped)
        if m:
            level = min(len(m.group(1)), 4)
            head = doc.add_heading("", level=level)
            _add_runs(head, _clean_inline(m.group(2)))
            for run in head.runs:
                run.font.name = BODY_FONT
                # sin el azul por defecto de Word: escrito judicial sobrio en negro
                run.font.color.rgb = RGBColor(0, 0, 0)
            if level == 1:
                head.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        if _BULLET.match(stripped):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, _clean_inline(_BULLET.sub("", stripped, count=1)))
            continue
        if _NUMBERED.match(stripped):
            p = doc.add_paragraph(style="List Number")
            _add_runs(p, _clean_inline(_NUMBERED.sub("", stripped, count=1)))
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _add_runs(p, _clean_inline(stripped))

    # Pie de página: recordatorio de la regla de la casa en lenguaje del oficio.
    footer_p = doc.sections[0].footer.paragraphs[0]
    footer_p.text = FOOTER_NOTE
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer_p.runs:
        run.font.size = Pt(8)
        run.font.name = BODY_FONT

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
